from openpyxl import load_workbook
import json,zipfile
p='賸餘款繳回簽辦單_重新排版.xlsx'
w=load_workbook(p)
for s in w:
 print('SHEET',s.title,s.max_row,s.max_column,'PRINT',s.print_area,'PAGE',s.page_setup,'MARGINS',s.page_margins)
 print('MERGES',str(s.merged_cells))
 print('COLS',[(k,v.width,v.hidden) for k,v in s.column_dimensions.items()])
 print('ROWS',[(k,v.height,v.hidden) for k,v in s.row_dimensions.items()])
 for row in s:
  for c in row:
   if c.value is not None: print(c.coordinate,repr(c.value),'style',c.style_id,'font',c.font.name,c.font.sz,'bold',c.font.b,'align',c.alignment,'border',c.border)
with zipfile.ZipFile(p) as z: print('FILES',z.namelist())